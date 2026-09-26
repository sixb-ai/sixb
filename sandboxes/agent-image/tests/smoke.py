"""Offline acceptance checks for the built image. Writes reviewable artifacts to a new directory."""

# Regression checks: replace scripts/python3 with an external symlink to the venv interpreter;
# Python 3.11 then loses venv discovery and the imports below fail. Remove sixb-render's OOXML
# package check; LibreOffice 7.4 accepts broken.pptx as plain text and the final negative test fails.
# Count pages with pypdf instead of pdfinfo, or drop its cleanup; the corrupt PDF check fails.

import hashlib
import json
import subprocess
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from extended import check_extended


def command(*args, ok=True):
    result = subprocess.run(args, capture_output=True, text=True, timeout=90)
    if ok and result.returncode != 0:
        raise AssertionError(f"{args[0]}: {result.stderr or result.stdout}")
    if not ok and result.returncode == 0:
        raise AssertionError(f"Expected failure: {args}")
    return result


root = Path(sys.argv[1]).resolve()
root.mkdir(parents=True, exist_ok=False)
assert command("bun", "-e", "console.log(6 * 7)").stdout.strip() == "42"
assert command("sqlite3", ":memory:", "select 6 * 7;").stdout.strip() == "42"
assert command("jq", "-n", "6 * 7").stdout.strip() == "42"

# Test imports, real arithmetic, and image output without a display server.
frame = pd.DataFrame({"item": ["Diagnosis", "Part", "Labor"], "cost": [350, 1600, 500]})
assert int(np.sum(frame["cost"])) == 2450
frame.to_csv(root / "costs.csv", index=False)
frame.plot.bar(x="item", y="cost", legend=False)
plt.ylabel("USD")
plt.tight_layout()
plt.savefig(root / "costs.png")
plt.close()

pdfmetrics.registerFont(TTFont("SixbSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
pdf = canvas.Canvas(str(root / "report.pdf"))
for number in (1, 2):
    pdf.setFont("SixbSans", 24)
    pdf.drawString(60, 760, "Cedar service report")
    pdf.setFont("SixbSans", 14)
    pdf.drawString(60, 710, "SC-2087 | RTU-9 | Response target: 75 minutes")
    pdf.drawString(60, 675, "Estimated cost: $2,450. Approval pending.")
    pdf.drawString(60, 640, "Font check: café, £, €")
    pdf.drawString(60, 70, f"Page {number} of 2")
    pdf.showPage()
pdf.save()

doc = Document()
doc.add_heading("Cedar service report", 0)
doc.add_paragraph("SC-2087: RTU-9. Estimated cost $2,450; approval pending.")
table = doc.add_table(rows=1, cols=2)
table.rows[0].cells[0].text = "Item"
table.rows[0].cells[1].text = "USD"
for item, cost in zip(frame["item"], frame["cost"]):
    row = table.add_row().cells
    row[0].text, row[1].text = item, str(cost)
doc.save(root / "report.docx")

book = Workbook()
sheet = book.active
sheet.append(["Item", "USD"])
for item, cost in zip(frame["item"], frame["cost"]):
    sheet.append([item, int(cost)])
sheet.append(["Total", "=SUM(B2:B4)"])
sheet.column_dimensions["A"].width = 24
sheet.column_dimensions["B"].width = 16
book.save(root / "costs.xlsx")
assert load_workbook(root / "costs.xlsx").active["B5"].value == "=SUM(B2:B4)"

js = root / "slides.cjs"
js.write_text('''
const pptxgen = require("/opt/sixb/js/node_modules/pptxgenjs")
const pptx = new pptxgen()
pptx.layout = "LAYOUT_WIDE"
const slide = pptx.addSlide()
slide.addText("Cedar service costs", {x:0.6, y:0.4, w:12, h:0.7, fontSize:32, color:"17324D"})
slide.addText("SC-2087 | Estimated total $2,450 | Approval pending", {x:0.6, y:1.2, w:12, h:0.5, fontSize:18})
slide.addChart(pptx.ChartType.bar, [{name:"USD", labels:["Diagnosis","Part","Labor"], values:[350,1600,500]}],
 {x:0.6, y:2.1, w:11.8, h:4.5, showLegend:false, showValue:true, catAxisLabelFontSize:16})
pptx.writeFile({fileName:process.argv[2]}).catch(e=>{console.error(e);process.exitCode=1})
''')
command("node", str(js), str(root / "slides.pptx"))
deck = Presentation(root / "slides.pptx")
assert len(deck.slides) == 1
charts = [shape.chart for shape in deck.slides[0].shapes if shape.has_chart]
assert len(charts) == 1 and list(charts[0].series[0].values) == [350, 1600, 500]
# Exercise python-pptx editing, not merely its importer.
deck.slides[0].shapes[0].text = "Cedar service costs"
deck.save(root / "edited.pptx")

results = []
for filename, pages, required in [
    ("report.pdf", 2, "2,450"),
    ("slides.pptx", 1, "2,450"),
    ("edited.pptx", 1, "2,450"),
    ("report.docx", 1, "2,450"),
    ("costs.xlsx", 1, "2450"),
]:
    output = root / f"{filename}-preview"
    manifest = json.loads(command("sixb-render", str(root / filename), "--output-dir", str(output)).stdout)
    assert manifest["totalPages"] == pages and manifest["complete"]
    text = "\n".join(page.extract_text() for page in PdfReader(manifest["pdf"]).pages)
    assert required in text, (filename, text)
    for path in manifest["pages"]:
        with Image.open(path) as preview:
            assert max(preview.size) <= 1600
            assert preview.convert("L").getextrema()[0] < 200
    results.append({"file": filename, "pages": pages, "textChecked": required})

# A rerender cannot silently reuse or overwrite a prior result.
previous = root / "report.pdf-preview" / "document.pdf"
digest = hashlib.sha256(previous.read_bytes()).hexdigest()
command("sixb-render", str(root / "report.pdf"), "--output-dir", str(previous.parent), ok=False)
assert hashlib.sha256(previous.read_bytes()).hexdigest() == digest

# Page limits fail explicitly, and deliberate partial previews identify their scope.
command("sixb-render", str(root / "report.pdf"), "--output-dir", str(root / "limited"), "--max-pages", "1", ok=False)
assert not (root / "limited").exists()
partial = json.loads(command("sixb-render", str(root / "report.pdf"), "--output-dir", str(root / "partial"), "--first-page", "2").stdout)
assert not partial["complete"] and len(partial["pages"]) == 1
broken = root / "broken.pptx"
broken.write_text("not a presentation")
command("sixb-render", str(broken), "--output-dir", str(root / "broken-preview"), ok=False)
assert not (root / "broken-preview").exists()
# An unreadable PDF is a reported failure, not a traceback, and the directory is free for a retry.
corrupt = root / "corrupt.pdf"
corrupt.write_text("not a pdf")
failed = command("sixb-render", str(corrupt), "--output-dir", str(root / "corrupt-preview"), ok=False)
assert failed.stderr.startswith("[SixbRender]") and "Traceback" not in failed.stderr, failed.stderr
assert not (root / "corrupt-preview").exists()

extended = check_extended(root)
summary = {"status": "passed", "artifacts": results, **extended}
(root / "results.json").write_text(json.dumps(summary, indent=2) + "\n")
print(json.dumps(summary, indent=2))
